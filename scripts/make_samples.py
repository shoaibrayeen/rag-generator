"""
Generate the binary sample documents (PDF, DOCX, scanned-style PDF, PNG) in samples/.

    pip install -r requirements-dev.txt
    python scripts/make_samples.py

The text PDF and DOCX exercise the normal extractors; `scanned_memo.pdf` and
`whiteboard.png` contain text only as pixels, so they exercise the Tesseract OCR path.
"""
from __future__ import annotations

from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "samples"

WARRANTY_PAGES = [
    "Northwind Robotics Limited Warranty\n\nThis warranty covers manufacturing defects in Northwind "
    "robot arms, mobile bases and end effectors. Coverage starts on the delivery date shown on the "
    "commercial invoice.",
    "Coverage periods\n\nDesktop and collaborative arms: 2 years. Industrial arms (Titan series): "
    "3 years. Mobile bases: 2 years. Accessories and end effectors: 1 year. Extended coverage of up "
    "to 5 years can be purchased within 90 days of delivery.",
    "Exclusions\n\nThe warranty does not cover damage caused by operation outside the rated payload, "
    "unauthorised firmware, liquid ingress, or use in explosive atmospheres. Consumables such as "
    "gripper pads and cables are excluded.",
    "Claims\n\nTo make a claim, open a ticket at support.northwind.example quoting the serial number. "
    "Northwind will respond within 2 business days and ship replacement parts within 10 business days "
    "for in-stock items.",
]

MEMO = ("INTERNAL MEMO - SCANNED COPY\n\nTo: All plant supervisors\nFrom: Operations\n"
        "Date: 3 March 2026\n\nThe Rotterdam assembly line will pause for\n"
        "scheduled maintenance from 9 to 11 April 2026.\nShift leads must complete the\n"
        "safety checklist form OPS-17 before restart.\nContact: ops@northwind.example")

WHITEBOARD = ("Sprint 42 goals\n- Ship Fleet Manager 1.3\n- Reduce Rover XL boot time to 20 seconds\n"
              "- Hire two firmware engineers\nDemo day: Friday 26 June 2026")


def make_text_pdf() -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(OUT / "warranty.pdf"), pagesize=A4)
    width, height = A4
    for i, page in enumerate(WARRANTY_PAGES, start=1):
        y = height - 3 * cm
        for line in _wrap(page, 90):
            c.drawString(2.5 * cm, y, line)
            y -= 0.6 * cm
        c.drawString(width / 2, 1.5 * cm, f"Page {i}")
        c.showPage()
    c.save()


def make_docx() -> None:
    import docx

    d = docx.Document()
    d.add_heading("Onboarding Guide for New Engineers", level=0)
    d.add_heading("Week 1", level=1)
    d.add_paragraph("Collect your laptop from IT, complete the security training in the LMS, and "
                    "pair with your onboarding buddy. Your buddy is assigned in the welcome email.")
    d.add_heading("Development environment", level=1)
    d.add_paragraph("All services run in Docker Compose. Clone the 'northwind/monorepo' repository, "
                    "run make bootstrap, and confirm that make test passes before your first PR.")
    d.add_heading("Code review", level=1)
    d.add_paragraph("Every change needs one approving review. Changes to the motion planner or "
                    "safety controller need two approvals, one from the Safety team.")
    d.add_heading("Useful contacts", level=1)
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Team", "Channel"
    for team, chan in [("Firmware", "#fw"), ("Motion planning", "#motion"), ("IT helpdesk", "#it-help")]:
        row = table.add_row().cells
        row[0].text, row[1].text = team, chan
    d.save(str(OUT / "onboarding.docx"))


def _render_text_image(text: str, size=(1400, 900)):
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 34)
    except OSError:
        font = ImageFont.load_default(size=34)
    y = 60
    for line in text.split("\n"):
        draw.text((70, y), line, fill="black", font=font)
        y += 52
    return img


def make_scanned_pdf_and_png() -> None:
    _render_text_image(MEMO).save(OUT / "scanned_memo.pdf", "PDF", resolution=150)
    _render_text_image(WHITEBOARD, size=(1400, 600)).save(OUT / "whiteboard.png")


def _wrap(text: str, width: int) -> list[str]:
    import textwrap

    out: list[str] = []
    for para in text.split("\n"):
        out.extend(textwrap.wrap(para, width) or [""])
    return out


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    make_text_pdf()
    make_docx()
    make_scanned_pdf_and_png()
    print("wrote:", ", ".join(sorted(p.name for p in OUT.iterdir())))
